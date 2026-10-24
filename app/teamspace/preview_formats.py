"""Bounded local document previews. No external converters, macros or formula evaluation.
Office parsing runs in a disposable resource-limited subprocess at the HTTP boundary.
"""
from __future__ import annotations
import base64,csv,datetime,io,itertools,json,math,re,sys,zipfile
from pathlib import Path
MAX_BYTES=20*1024*1024
INLINE={'.pdf':('pdf','application/pdf'),'.png':('image','image/png'),'.jpg':('image','image/jpeg'),'.jpeg':('image','image/jpeg'),'.webp':('image','image/webp'),'.gif':('image','image/gif'),'.mp3':('audio','audio/mpeg'),'.wav':('audio','audio/wav'),'.ogg':('audio','audio/ogg'),'.mp4':('video','video/mp4'),'.webm':('video','video/webm')}
TEXT={'.md','.markdown','.txt','.json','.jsonl','.csv','.tsv','.yaml','.yml','.toml','.py','.js','.ts','.sh','.ps1','.html','.css','.xml','.ini','.cfg','.log','.sql','.rst','.bat'}

def checked_zip(raw):
 z=zipfile.ZipFile(io.BytesIO(raw));infos=z.infolist()
 if len(infos)>3000 or sum(x.file_size for x in infos)>60*1024*1024:raise ValueError('Office压缩包展开过大')
 for x in infos:
  if x.flag_bits&1 or x.file_size>20*1024*1024 or x.file_size>max(x.compress_size,1)*1000:raise ValueError('Office包加密或超过展开限制')
  if x.filename.endswith(('.xml','.rels')):
   xml=z.read(x)
   if b'<!DOCTYPE' in xml.upper() or b'<!ENTITY' in xml.upper():raise ValueError('不支持含外部实体的XML')
 return z

def decode(raw):
 for encoding in ('utf-8-sig','utf-16' if raw[:2] in (b'\xff\xfe',b'\xfe\xff') else 'gb18030'):
  try:
   text=raw.decode(encoding)
   if '\x00' not in text:return text
  except UnicodeError:pass
 raise ValueError('文本编码不支持；请另存UTF-8')

def cell_text(value,fmt=''):
 if value is None:return ''
 if isinstance(value,(datetime.datetime,datetime.date,datetime.time)):return value.isoformat(sep=' ') if isinstance(value,datetime.datetime) else value.isoformat()
 if isinstance(value,bool):return 'TRUE' if value else 'FALSE'
 if isinstance(value,(int,float)) and not isinstance(value,bool):
  if isinstance(value,float) and not math.isfinite(value):return str(value)
  if '%' in fmt:
   decimals=len(re.search(r'\.([0#]+)',fmt).group(1)) if re.search(r'\.([0#]+)',fmt) else 0
   return f'{value*100:.{min(decimals,8)}f}%'
  if re.fullmatch(r'0\.0+',fmt):return f'{value:.{min(len(fmt)-2,8)}f}'
 return str(value)[:8000]

def preview(raw,filename,sheet=0,row_offset=0,col_offset=0):
 if not raw or len(raw)>MAX_BYTES:raise ValueError('文件为空或超过20MB预览限制')
 ext=Path(filename).suffix.lower()
 if ext in INLINE:
  kind,mime=INLINE[ext]
  if kind=='pdf' and not raw.startswith(b'%PDF-'):raise ValueError('不是有效PDF')
  if kind=='image':
   from PIL import Image
   with Image.open(io.BytesIO(raw)) as im:
    if im.width*im.height>40_000_000:raise ValueError('图片像素过大')
    im.verify()
  return {'format':kind,'mime':mime,'bytes':len(raw),'original':True}
 if ext in TEXT or Path(filename).name in ('LICENSE','NOTICE'):
  text=decode(raw)
  if ext=='.json':
   if len(text)<2_000_000:
    try:text=json.dumps(json.loads(text),ensure_ascii=False,indent=2)
    except (ValueError,RecursionError):pass
  if ext in ('.csv','.tsv'):
   reader=csv.reader(io.StringIO(text),delimiter='\t' if ext=='.tsv' else ',');rows=list(itertools.islice(reader,row_offset,row_offset+101));more=len(rows)>100
   return {'format':'table','sheet':0,'sheets':['数据'],'rows':[[str(c)[:4000] for c in row[col_offset:col_offset+40]] for row in rows[:100]],'row_offset':row_offset,'col_offset':col_offset,'has_more_rows':more,'has_more_cols':any(len(row)>col_offset+40 for row in rows),'formula_policy':'仅显示，不计算公式'}
  start=row_offset;size=50000
  return {'format':'markdown' if ext in ('.md','.markdown') else 'text','content':text[start:start+size],'language':ext.lstrip('.'),'offset':start,'has_more':start+size<len(text),'next_offset':start+size,'total_characters':len(text)}
 if ext=='.docx':
  with checked_zip(raw):pass
  from docx import Document
  from docx.table import Table
  doc=Document(io.BytesIO(raw));blocks=[];images=0;image_bytes=0;chars=0;limited=False
  for block in doc.iter_inner_content():
   if len(blocks)>=1200 or chars>200000:limited=True;break
   if isinstance(block,Table):
    rows=[]
    for row in itertools.islice(block.rows,150):
     vals=[cell.text[:6000] for cell in itertools.islice(row.cells,30)];chars+=sum(map(len,vals));rows.append(vals)
    blocks.append({'type':'table','rows':rows,'truncated':len(block.rows)>150})
   else:
    style=block.style.name if block.style else '';text=block.text;chars+=len(text)
    level=int(re.search(r'(\d+)',style).group(1)) if style.startswith('Heading') and re.search(r'(\d+)',style) else 0
    runs=[{'text':run.text,'bold':bool(run.bold),'italic':bool(run.italic),'underline':bool(run.underline)} for run in block.runs]
    blocks.append({'type':'paragraph','text':text[:20000],'runs':runs[:1000],'heading':min(level,6),'list':style.startswith('List'),'align':str(block.alignment)})
    for b in block._p.xpath('.//a:blip'):
     rid=b.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed');part=doc.part.related_parts.get(rid)
     if part and part.content_type in ('image/png','image/jpeg','image/gif','image/webp') and images<20 and len(part.blob)<=1024*1024 and image_bytes+len(part.blob)<=2*1024*1024:
      from PIL import Image
      try:
       with Image.open(io.BytesIO(part.blob)) as im:
        if im.width*im.height>20_000_000:continue
        im.verify()
       blocks.append({'type':'image','mime':part.content_type,'base64':base64.b64encode(part.blob).decode()});images+=1;image_bytes+=len(part.blob)
      except Exception:continue
  return {'format':'word','blocks':blocks,'truncated':limited,'notice':'常规内容预览；显示正文、表格和内嵌图片，不运行宏，不保证与Word逐页排版一致。'}
 if ext in ('.xlsx','.xlsm','.xls'):
  if sheet<0 or row_offset<0 or col_offset<0 or row_offset>100000 or col_offset>1000:raise ValueError('表格分页参数超限')
  if ext=='.xls':
   import xlrd
   book=xlrd.open_workbook(file_contents=raw,on_demand=True)
   try:
    names=book.sheet_names()
    if sheet>=len(names):raise ValueError('工作表不存在')
    ws=book.sheet_by_index(sheet);rows=[]
    for i in range(row_offset,min(row_offset+100,ws.nrows)):
     vals=[]
     for j in range(col_offset,min(col_offset+40,ws.ncols)):
      c=ws.cell(i,j);val=c.value
      if c.ctype==xlrd.XL_CELL_DATE:
       try:val=xlrd.xldate_as_datetime(val,book.datemode)
       except ValueError:pass
      vals.append(cell_text(val))
     rows.append(vals)
    return {'format':'table','sheets':names,'sheet':sheet,'rows':rows,'row_offset':row_offset,'col_offset':col_offset,'total_rows':ws.nrows,'total_cols':ws.ncols,'has_more_rows':row_offset+100<ws.nrows,'has_more_cols':col_offset+40<ws.ncols,'formula_policy':'显示保存的值，不重新计算公式'}
   finally:book.release_resources()
  with checked_zip(raw):pass
  from openpyxl import load_workbook
  book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False);formulas=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
  try:
   names=book.sheetnames
   if sheet>=len(names):raise ValueError('工作表不存在')
   ws=book[names[sheet]];fs=formulas[names[sheet]];nr=ws.max_row or 0;nc=ws.max_column or 0;rows=[]
   kw={'min_row':row_offset+1,'max_row':min(row_offset+100,nr),'min_col':col_offset+1,'max_col':min(col_offset+40,nc)}
   if nr>row_offset and nc>col_offset:
    for values,raws in zip(ws.iter_rows(**kw),fs.iter_rows(**kw)):
     rows.append([cell_text(v.value if v.value is not None else f.value,getattr(v,'number_format','')) for v,f in zip(values,raws)])
   return {'format':'table','sheets':names,'sheet':sheet,'rows':rows,'row_offset':row_offset,'col_offset':col_offset,'total_rows':nr,'total_cols':nc,'has_more_rows':row_offset+100<nr,'has_more_cols':col_offset+40<nc,'formula_policy':'优先显示保存的计算值；没有缓存时显示公式本身，不执行公式、宏或外部链接'}
  finally:book.close();formulas.close()
 return {'format':'unsupported','notice':'此格式暂不支持内嵌预览，可下载原件。旧版 .doc 请转换为 .docx；不向第三方上传文件。'}

if __name__=='__main__':
 try:
  import resource
  resource.setrlimit(resource.RLIMIT_CPU,(20,25));resource.setrlimit(resource.RLIMIT_AS,(1200*1024*1024,1200*1024*1024))
  data=preview(Path(sys.argv[1]).read_bytes(),sys.argv[2],*map(int,sys.argv[3:6]));print(json.dumps(data,ensure_ascii=False,allow_nan=False))
 except Exception as e:print(json.dumps({'format':'error','notice':str(e)[:200] if isinstance(e,ValueError) else '预览失败（'+type(e).__name__+'），原件保留'},ensure_ascii=False));sys.exit(0)
