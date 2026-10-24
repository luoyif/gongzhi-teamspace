"""Bounded native professional outputs. No uploaded code, macros or shell is executed."""
from __future__ import annotations
import ast,hashlib,io,json,re,zipfile
from pathlib import PurePosixPath
from .store import Problem,dump,require_text

CATALOG=[
 {'id':'office.documents','name':'Word文档交付','description':'将已核对内容生成真实DOCX，含标题、段落、表格和来源；保存前确认，不覆盖原件或执行宏。','tools':['office_docx_create'],'write':True},
 {'id':'engineering.deliver','name':'代码交付与静态检查','description':'检查Python语法并保存代码ZIP；不执行代码、不安装依赖、不访问宿主源码或发布生产。','tools':['python_syntax_check','code_bundle_save'],'write':True},
]

def text(value,label,limit=12000,empty=True):
 value=require_text(value,label,limit,empty)
 if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',value):raise Problem(422,label+'含不支持的控制字符')
 return value

def file_name(name,extension):
 if not isinstance(name,str) or not re.fullmatch(r'[\w .\-\u3400-\u9fff]{1,110}\.'+extension,name) or '..' in name:raise Problem(422,'请输入安全的 '+extension+' 文件名，不接受路径')
 return name

def docx_bytes(title,sections,subtitle='',sources=()):
 from docx import Document
 from docx.shared import Cm,Pt,RGBColor
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 title=text(title,'文档标题',160,False);subtitle=text(subtitle,'副标题',500)
 if not isinstance(sections,list) or not 1<=len(sections)<=40:raise Problem(422,'文档需要1–40个章节')
 if len(dump(sections))>100000:raise Problem(413,'文档正文过大，请分卷')
 d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=Cm(2.0);sec.left_margin=sec.right_margin=Cm(2.1)
 for style in ('Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet'):
  s=d.styles[style];s.font.name='Calibri';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');s.font.color.rgb=RGBColor.from_string('202124')
 normal=d.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.18
 d.styles['Title'].font.size=Pt(24);d.styles['Heading 1'].font.size=Pt(15);d.styles['Heading 1'].paragraph_format.space_before=Pt(16);d.styles['Heading 1'].paragraph_format.keep_with_next=True
 d.add_heading(title,0)
 if subtitle:d.add_paragraph(subtitle,'Subtitle')
 plain=[title,subtitle]
 for section in sections:
  if not isinstance(section,dict) or set(section)-{'heading','paragraphs','bullets','table'}:raise Problem(422,'章节仅支持heading、paragraphs、bullets、table')
  heading=text(section.get('heading',''),'章节标题',200)
  if heading:d.add_heading(heading,1);plain.append(heading)
  for key,style in (('paragraphs',None),('bullets','List Bullet')):
   values=section.get(key,[])
   if not isinstance(values,list) or len(values)>60:raise Problem(422,'段落或条目数量超限')
   for value in values:
    value=text(value,'段落',12000);d.add_paragraph(value,style);plain.append(value)
  table=section.get('table')
  if table is not None:
   if not isinstance(table,dict) or set(table)-{'headers','rows'}:raise Problem(422,'表格只接受headers和rows')
   headers=table.get('headers',[]);rows=table.get('rows',[])
   if not isinstance(headers,list) or not 1<=len(headers)<=8 or not isinstance(rows,list) or len(rows)>100:raise Problem(422,'表格限制8列、100行')
   data=[headers]+rows
   if any(not isinstance(row,list) or len(row)!=len(headers) for row in data):raise Problem(422,'表格每行列数须一致')
   t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
   for cell in t.rows[0].cells:cell.width=Cm(16.8/len(headers))
   repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');t.rows[0]._tr.get_or_add_trPr().append(repeat)
   for ri,row in enumerate(data):
    cells=t.rows[0].cells if ri==0 else t.add_row().cells;plain.append(' | '.join(str(v) for v in row))
    for ci,value in enumerate(row):
     value=text(str(value),'表格单元格',4000);cells[ci].text=value
     for p in cells[ci].paragraphs:
      p.paragraph_format.space_after=Pt(4)
      for run in p.runs:run.font.size=Pt(9);run.bold=ri==0
     if ri==0:
      shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEF2F5');cells[ci]._tc.get_or_add_tcPr().append(shade)
 if sources:
  d.add_heading('来源与版本',1)
  for source in sources:
   value=source['title']+' · '+source['document_id']+' · v'+str(source['revision']);d.add_paragraph(value);plain.append(value)
 footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('共知 · 待人工验收   ')
 page=OxmlElement('w:fldSimple');page.set(qn('w:instr'),'PAGE');footer._p.append(page)
 d.core_properties.author='共知团队';d.core_properties.title=title;d.core_properties.subject='待人工验收的项目产物'
 buf=io.BytesIO();d.save(buf)
 return buf.getvalue(),'\n\n'.join(x for x in plain if x)

def code_zip(files):
 if not isinstance(files,list) or not 1<=len(files)<=30:raise Problem(422,'代码包需1–30个文本文件')
 out=io.BytesIO();names=set();plain=[];total=0
 with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
  for entry in files:
   if not isinstance(entry,dict) or set(entry)!={'path','content'}:raise Problem(422,'文件需path和content')
   name=entry['path']
   if not isinstance(name,str) or len(name)>180:raise Problem(422,'代码文件路径无效')
   p=PurePosixPath(name)
   if p.is_absolute() or p.as_posix()!=name or '\\' in name or ':' in name or any(x in ('.','..','') or x.startswith('.') for x in name.split('/')) or name.casefold() in names or any(ord(x)<32 for x in name):raise Problem(422,'代码文件路径不安全或重名')
   if p.suffix.lower() not in {'.py','.js','.ts','.tsx','.jsx','.html','.css','.json','.yaml','.yml','.md','.txt','.csv','.sql','.toml','.xml'}:raise Problem(422,'只支持代码/配置/文档文本，不支持二进制程序或密钥文件')
   content=text(entry['content'],'代码文件',70000);total+=len(content.encode())
   if total>250000:raise Problem(413,'代码包文本超过250KB')
   if p.suffix.lower()=='.json':
    try:json.loads(content)
    except ValueError:raise Problem(422,'JSON文件语法错误：'+name) from None
   names.add(name.casefold());z.writestr(name,content);plain.append('## '+name+'\n\n'+content)
 return out.getvalue(),'\n\n'.join(plain)

def syntax_check(code):
 code=text(code,'Python源码',40000,False)
 try:tree=ast.parse(code,mode='exec');return {'syntax_valid':True,'top_level_statements':len(tree.body),'executed':False,'functional_tests_run':False}
 except (SyntaxError,ValueError,RecursionError) as e:return {'syntax_valid':False,'line':getattr(e,'lineno',None),'column':getattr(e,'offset',None),'message':getattr(e,'msg',type(e).__name__),'executed':False,'functional_tests_run':False}

def store_output(ops,c,filename,raw,body,sources):
 digest=hashlib.sha256(raw).hexdigest();relative='blobs/'+digest;path=ops.store.data_dir/relative
 if not path.exists():path.write_bytes(raw);path.chmod(0o600)
 data={**ops.provenance(),'document_kind':'file','filename':filename,'blob_path':relative,'sha256':digest,'bytes':len(raw),'source_documents':sources,'origin':'approved_professional_output','scripts_executed':False}
 id=ops.store.put_object(c,ops.p,'document',filename,body[:500000],ops.actor,data,'ready');ops.store.index(c,id,ops.p,[('待验收专业产物',body[:500000])],filename)
 return {'id':id,'title':filename,'revision':1,'status':'ready','adopted':False,'sha256':digest,'bytes':len(raw),'download_path':'/api/team/documents/'+id+'/download','executed':False,'visual_review':'not_automatic'}

def functions(ops):
 async def python_syntax_check(code:str)->str:
  """Parse Python syntax without executing code. A pass is NOT functional testing or business correctness."""
  with ops.store.connection() as c:ops.authorize(c)
  return dump(syntax_check(code))
 async def code_bundle_save(filename:str,files:list[dict])->str:
  """AFTER approval save a ZIP of text source files. Each file: {path,content}. Never runs, deploys, installs or overwrites host source. Max30 files/250KB."""
  file_name(filename,'zip');raw,body=code_zip(files)
  return ops.mutation('code_bundle_save',{'filename':filename,'files':files},lambda c:store_output(ops,c,filename,raw,body,[]))
 async def office_docx_create(filename:str,title:str,sections:list[dict],source_document_ids:list[str],subtitle:str='')->str:
  """AFTER approval create an actual Word DOCX. Sections: heading,paragraphs:list[str],bullets:list[str],table:{headers:list[str],rows:list[list[str]]}. Sources are current project document IDs. Does not overwrite originals or fabricate verification."""
  file_name(filename,'docx')
  if not isinstance(source_document_ids,list) or len(source_document_ids)>20 or any(not isinstance(x,str) for x in source_document_ids):raise Problem(422,'来源文件ID无效')
  def apply(c):
   refs=[]
   for id in dict.fromkeys(source_document_ids):
    d=ops.registry.resolve_document(ops.p,ops.actor,id,c);refs.append({'document_id':id,'title':d['title'],'revision':d['revision']})
   raw,body=docx_bytes(title,sections,subtitle,refs);return store_output(ops,c,filename,raw,body,refs)
  return ops.mutation('office_docx_create',{'filename':filename,'title':title,'sections':sections,'source_document_ids':source_document_ids,'subtitle':subtitle},apply)
 return {f.__name__:f for f in (python_syntax_check,code_bundle_save,office_docx_create)}
