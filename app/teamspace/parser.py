"""Bounded file parsing subprocess; no macros, shell commands, or OCR."""
from __future__ import annotations
import json
import sys
import zipfile
from pathlib import Path


def parse(path: Path, filename: str):
    suffix=Path(filename).suffix.lower()
    if suffix in {".md",".txt",".csv",".json",".log"}:
        raw=path.read_bytes()
        try: text=raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try: text=raw.decode("gb18030")
            except UnicodeDecodeError: return "unindexed",[],"暂不支持此文本编码，请另存为UTF-8"
        if "\x00" in text: return "unindexed",[],"文件包含二进制内容"
        if len(text)>500000: return "unindexed",[],"文本超过50万字符；原件已保存，请分成较小文件"
        return "ready",[("正文",text)],None
    if suffix in {".xlsx",".xlsm",".xls"}:
        from .preview_formats import preview
        raw=path.read_bytes();first=preview(raw,filename);segments=[];size=0
        if len(first.get("sheets",[]))>30:return "unindexed",[],"工作表超过30张，请拆分后索引"
        for sheet,name in enumerate(first.get("sheets",[])):
            offset=0
            while True:
                page=preview(raw,filename,sheet,offset,0)
                if page.get("has_more_cols"):return "unindexed",[],"表格超过40列；可以预览，请拆分后建立完整文本索引"
                text="\n".join(" | ".join(row) for row in page["rows"]);size+=len(text)
                if size>500000 or offset>=10000:return "unindexed",[],"表格超过索引上限；原件仍可预览"
                if text.strip():segments.append((name+" · 行"+str(offset+1),text))
                if not page.get("has_more_rows"):break
                offset+=100
        return ("ready",segments,None) if segments else ("unindexed",[],"工作簿没有可读内容")
    if suffix==".pdf":
        from pypdf import PdfReader
        pdf=PdfReader(path)
        if pdf.is_encrypted: return "unindexed",[],"加密PDF已保存，请上传可读取版本"
        if len(pdf.pages)>150: return "unindexed",[],"PDF超过150页；原件已保存，请拆分后索引"
        segments=[]; length=0
        for i,page in enumerate(pdf.pages):
            text=page.extract_text() or ""; length+=len(text)
            if length>500000: return "unindexed",[],"提取文本超过50万字符，请拆分"
            if text.strip(): segments.append((f"第{i+1}页",text))
        if not segments: return "unindexed",[],"未找到文字层；扫描件需要OCR，原件已保存"
        return "ready",segments,None
    if suffix==".docx":
        with zipfile.ZipFile(path) as z:
            if sum(x.file_size for x in z.infolist())>60*1024*1024 or len(z.infolist())>3000:
                return "unindexed",[],"DOCX展开内容过大，原件已保存"
        from docx import Document
        doc=Document(path)
        lines=[p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows: lines.append(" | ".join(c.text for c in row.cells))
        text="\n\n".join(lines)
        if len(text)>500000: return "unindexed",[],"提取文本超过50万字符，请拆分"
        return ("ready",[("正文及表格",text)],None) if text.strip() else ("unindexed",[],"未提取到正文")
    return "unindexed",[],"此格式只保存原件，尚未建立文本索引"


if __name__=="__main__":
    try:
        try:
            import resource
            resource.setrlimit(resource.RLIMIT_CPU,(35,40))
            resource.setrlimit(resource.RLIMIT_AS,(1200*1024*1024,1200*1024*1024))
        except (ImportError,ValueError,OSError): pass
        status,segments,error=parse(Path(sys.argv[1]),sys.argv[2])
        print(json.dumps({"status":status,"segments":segments,"error":error},ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"status":"failed","segments":[],"error":f"解析失败（{type(exc).__name__}）；原件已保存"},ensure_ascii=False))
